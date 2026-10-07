"""docx / xlsx / xls / pptx / csv 解析单测（老仓库 test_parsing_office.py 离线用例对拍）。

上传 API 集成用例随 2.4 迁入；本文件 fixture 现场生成文件 bytes，不提交二进制。
xls 没有现成的生成库，happy path 由集成阶段真实文件验证；单测覆盖损坏拒收 + 白名单。
"""

from io import BytesIO

import pytest

from app.core.errors import DocumentParseError
from app.pipeline.parsing import (
    EMPTY_HINT,
    SUPPORTED_FORMATS_HINT,
    SUPPORTED_SUFFIXES,
    parse_document,
    parse_document_pages,
)


# ---------------------------------------------------------------------------
# 文件构造（现场生成，不留二进制文件）
# ---------------------------------------------------------------------------


def make_docx() -> bytes:
    import docx as docx_lib

    document = docx_lib.Document()
    document.add_heading("项目说明", level=1)
    document.add_paragraph("这是正文段落。")
    document.add_heading("数据来源", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "投影仪"
    table.cell(1, 1).text = "3"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_tricky_cells() -> bytes:
    """表格单元格里带竖线与换行，验证 Markdown 表格不串列。"""
    import docx as docx_lib

    document = docx_lib.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列一"
    table.cell(0, 1).text = "列二"
    first = table.cell(1, 0)
    first.paragraphs[0].text = "第一行"
    first.add_paragraph("第二行")
    table.cell(1, 1).text = "a|b"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_empty_xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()  # 只有默认的空 sheet
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide1 = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide1.shapes.title.text = "第一页标题"
    slide1.placeholders[1].text = "第一页正文"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide2.shapes.title.text = "第二页标题"
    slide2.notes_slide.notes_text_frame.text = "记得核对预算"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_empty_pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()  # 零页演示文稿
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------


def test_docx_extracts_headings_paragraph_and_table_in_order():
    text = parse_document("a.docx", make_docx())
    assert "# 项目说明" in text
    assert "这是正文段落。" in text
    assert "## 数据来源" in text
    assert "| 名称 | 数量 |" in text
    assert "| --- | --- |" in text
    assert "| 投影仪 | 3 |" in text
    # 保序：标题 1 → 段落 → 标题 2 → 表格，按文档流
    order = [text.index(part) for part in ("# 项目说明", "这是正文段落。", "## 数据来源", "| 名称 |")]
    assert order == sorted(order)


def test_docx_table_cell_escapes_pipe_and_newline():
    text = parse_document("a.docx", make_docx_with_tricky_cells())
    assert "a\\|b" in text
    assert "第一行 第二行" in text  # 单元格内换行拍平成空格
    # 整张表不串列：表头 + 分隔行 + 数据行，正好三条竖线行
    pipe_lines = [line for line in text.splitlines() if line.startswith("|")]
    assert len(pipe_lines) == 3


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


def test_xlsx_two_sheets_with_headers_and_rows():
    data = make_xlsx(
        {
            "设备": [["名称", "金额"], ["投影仪", 3], ["白板", 2]],
            "人员": [["姓名"], ["张三"]],
        }
    )
    text = parse_document("a.xlsx", data)
    assert "## 设备" in text
    assert "| 名称 | 金额 |" in text
    assert "| 投影仪 | 3 |" in text
    assert "| 白板 | 2 |" in text
    assert "## 人员" in text
    assert "| 张三 |" in text


def test_xlsx_truncates_over_5000_rows():
    rows = [["列A", "列B"]] + [[f"r{i}", i] for i in range(5001)]  # 共 5002 行
    text = parse_document("a.xlsx", make_xlsx({"数据": rows}))
    assert "> （Sheet「数据」超过 5000 行，已截断）" in text
    # 表格保留 5000 行（首行是表头），加一条分隔线，共 5001 条竖线行
    pipe_lines = [line for line in text.splitlines() if line.startswith("|")]
    assert len(pipe_lines) == 5001


def test_xlsx_all_sheets_empty_reports_empty():
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document("a.xlsx", make_empty_xlsx())
    assert str(excinfo.value) == EMPTY_HINT


# ---------------------------------------------------------------------------
# xls（无生成库：只测白名单 + 损坏拒收）
# ---------------------------------------------------------------------------


def test_xls_suffix_whitelisted_and_broken_bytes_rejected_readably():
    assert ".xls" in SUPPORTED_SUFFIXES
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document("a.xls", b"not a real xls")
    message = str(excinfo.value)
    assert "Excel 97-2003 解析失败" in message
    assert "暂不支持" not in message


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------


def test_pptx_extracts_slides_and_notes():
    text = parse_document("a.pptx", make_pptx())
    assert "## 幻灯片 1" in text
    assert "第一页标题" in text
    assert "第一页正文" in text
    assert "## 幻灯片 2" in text
    assert "第二页标题" in text
    assert "**备注：**记得核对预算" in text


def test_pptx_all_empty_reports_empty():
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document("a.pptx", make_empty_pptx())
    assert str(excinfo.value) == EMPTY_HINT


# ---------------------------------------------------------------------------
# csv
# ---------------------------------------------------------------------------


def test_csv_basic_table():
    text = parse_document("a.csv", "名称,金额\n投影仪,3\n白板,2".encode("utf-8"))
    assert text == "| 名称 | 金额 |\n| --- | --- |\n| 投影仪 | 3 |\n| 白板 | 2 |"


def test_csv_gbk_encoded_with_crlf():
    text = parse_document("a.csv", "名称,金额\r\n投影仪,3".encode("gbk"))
    assert "| 名称 | 金额 |" in text
    assert "| 投影仪 | 3 |" in text


def test_csv_single_line_without_comma_is_plain_text():
    assert parse_document("a.csv", "只有一行没有逗号".encode("utf-8")) == "只有一行没有逗号"


def test_csv_empty_reports_empty():
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document("a.csv", b"")
    assert str(excinfo.value) == EMPTY_HINT


# ---------------------------------------------------------------------------
# 边界：空 bytes / 损坏 bytes 一律可读错误，不是 500
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("filename", ["a.docx", "a.xlsx", "a.xls", "a.pptx", "a.csv"])
def test_office_empty_bytes_reports_empty(filename):
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document(filename, b"")
    assert str(excinfo.value) == EMPTY_HINT


@pytest.mark.parametrize(
    ("filename", "prefix"),
    [
        ("a.docx", "Word 文档解析失败"),
        ("a.xlsx", "Excel 解析失败"),
        ("a.xls", "Excel 97-2003 解析失败"),
        ("a.pptx", "PPT 解析失败"),
        # csv 没法用垃圾 bytes 触发解析失败：任何 bytes 都能按 latin-1 解码并拆成行
    ],
)
def test_office_broken_bytes_reports_readable_error(filename, prefix):
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document(filename, b"na\x00me,not a real " + filename.encode())
    assert str(excinfo.value).startswith(prefix)


def test_office_pages_always_empty():
    """新格式没有页边界信息，pages 一律空列表（清洗的页眉页脚检测自然跳过）。"""
    cases = [
        ("a.docx", make_docx()),
        ("a.xlsx", make_xlsx({"数据": [["名称"], ["投影仪"]]})),
        ("a.csv", "名称,金额\n投影仪,3".encode("utf-8")),
        ("a.pptx", make_pptx()),
    ]
    for filename, data in cases:
        text, pages = parse_document_pages(filename, data)
        assert text.strip(), filename
        assert pages == [], filename


def test_unsupported_suffix_message_lists_all_formats():
    """不支持格式的提示要把**当前全部**可上传格式列全（阶段 3 起含六种图片）。

    2026-09-18 阶段 3 改动（见 S3-验收报告「登记偏差 P-登记-回归检查」）：本断言原先钉的是
    九格式列表（阶段 2 口径）。新增六种图片格式后该文案必须同步变长，否则用户看到的提示
    与实际支持能力不符（正是条目 3.3 要消除的误导）。属**断言随白名单扩展同步**，
    不是行为变更：九种原格式的解析结果、状态码、报错语义均未动。
    现在同一条断言直接读 `SUPPORTED_FORMATS_HINT`，避免以后每加一种格式都要手抄文案。
    """
    with pytest.raises(DocumentParseError) as excinfo:
        parse_document("a.exe", b"whatever")
    message = str(excinfo.value)
    assert "暂不支持" in message
    assert SUPPORTED_FORMATS_HINT in message
    # 六种图片后缀必须在提示里可见（条目 3.3 的验收点）
    for suffix in (".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".webp"):
        assert suffix.lstrip(".") in message
